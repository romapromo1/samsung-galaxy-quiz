import openpyxl
import json
import os
import sys

EXCEL_PATH = r"C:\Users\RocketPC\Downloads\Telegram Desktop\Samsung_NSW26_вопросы для квиза.xlsx"
OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src", "data")
OUTPUT_PATH = os.path.join(OUTPUT_DIR, "questions.json")

os.makedirs(OUTPUT_DIR, exist_ok=True)

print(f"Loading Excel file from: {EXCEL_PATH}")
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb['Вопросы']

norm_map = {
    'А': 'A', 'В': 'B', 'С': 'C',
    'A': 'A', 'B': 'B', 'C': 'C',
    'а': 'A', 'в': 'B', 'с': 'C',
    'a': 'A', 'b': 'B', 'c': 'C'
}

questions = []
for r in range(2, ws.max_row + 1):
    q_num = ws.cell(r, 1).value
    q_text = ws.cell(r, 2).value
    opt_a = ws.cell(r, 3).value
    opt_b = ws.cell(r, 4).value
    opt_c = ws.cell(r, 5).value
    ans_raw = str(ws.cell(r, 6).value or '').strip()

    if not q_text or not str(q_text).strip():
        continue

    ans = norm_map.get(ans_raw)
    if not ans:
        print(f"Warning: row {r} has unrecognized answer: '{ans_raw}'")
        continue

    options = [
        str(opt_a).strip() if opt_a is not None else '',
        str(opt_b).strip() if opt_b is not None else '',
        str(opt_c).strip() if opt_c is not None else ''
    ]

    correct_idx = 0 if ans == 'A' else (1 if ans == 'B' else 2)

    questions.append({
        'id': int(q_num) if q_num is not None else len(questions) + 1,
        'question': str(q_text).strip(),
        'options': options,
        'correctIndex': correct_idx,
        'correctAnswer': options[correct_idx]
    })

print(f"Successfully processed {len(questions)} questions.")

with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
    json.dump(questions, f, ensure_ascii=False, indent=2)

print(f"Saved to {OUTPUT_PATH}")
